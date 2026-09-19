"""Документы: PDF по страницам, PNG слайдов, vision, DOCX, TXT, деление постов."""

from __future__ import annotations

import asyncio
import base64
import logging
import os
import re
import shutil
import subprocess
from pathlib import Path
from typing import List, Optional, Tuple

logger = logging.getLogger(__name__)

_POST_SPLIT_RE = re.compile(r"(?:\n---+\n|\n\*\*\*+\n|\n{4,})")


def read_text_file(path: str) -> str:
    raw = Path(path).read_bytes()
    for enc in ("utf-8", "cp1251"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_docx(path: str) -> str:
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if (p.text or "").strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def pdftotext_pages(path: str) -> List[Tuple[int, str]]:
    if not shutil.which("pdftotext"):
        logger.warning("pdftotext не найден")
        return []
    proc = subprocess.run(
        ["pdftotext", "-layout", "-enc", "UTF-8", path, "-"],
        capture_output=True,
        timeout=180,
        check=False,
    )
    if proc.returncode != 0:
        logger.warning("pdftotext rc=%s", proc.returncode)
        return []
    text = proc.stdout.decode("utf-8", errors="replace")
    pages = text.split("\f")
    out: List[Tuple[int, str]] = []
    for i, page in enumerate(pages, start=1):
        t = page.strip()
        if t:
            out.append((i, t))
        else:
            out.append((i, ""))
    if out and not out[-1][1] and len(out) > 1:
        out = out[:-1]
    return out


def pdftoppm_png(path: str, dest_dir: str, *, scale_to: int = 1280) -> List[Path]:
    dest = Path(dest_dir)
    dest.mkdir(parents=True, exist_ok=True)
    prefix = dest / "page"
    if not shutil.which("pdftoppm"):
        return []
    proc = subprocess.run(
        ["pdftoppm", "-png", "-scale-to", str(int(scale_to)), path, str(prefix)],
        capture_output=True,
        timeout=300,
        check=False,
    )
    if proc.returncode != 0:
        logger.warning("pdftoppm rc=%s", proc.returncode)
        return []
    files = sorted(dest.glob("page-*.png")) or sorted(dest.glob("page*.png"))
    return files


def split_posts(text: str, *, min_chars: int = 200) -> List[str]:
    parts = _POST_SPLIT_RE.split(text or "")
    return [p.strip() for p in parts if len(p.strip()) >= min_chars]


async def vision_page(
    png_path: Path,
    openai_client,
    user_id: int,
    *,
    model: str = "gpt-4o-mini",
) -> str:
    raw = png_path.read_bytes()
    b64 = base64.b64encode(raw).decode("ascii")
    prompt = "перепиши текст слайда и кратко опиши схему или картинку"
    text = await openai_client.describe_image(b64, user_id, prompt=prompt)
    return (text or "").strip()


async def extract_document(
    path: str,
    *,
    kind: str,
    dest_dir: str,
    openai_client,
    user_id: int,
) -> dict:
    """
    Возвращает {
      pages: [{page, text, vision: bool}],
      method: str,
      slides: [png paths],
    }
    """
    from config import config

    ext = Path(path).suffix.lower()
    dest = Path(dest_dir)
    dest.mkdir(parents=True, exist_ok=True)
    if ext in {".txt", ".md"}:
        text = read_text_file(path)
        return {
            "pages": [{"page": 1, "text": text, "vision": False}],
            "method": "txt",
            "slides": [],
            "full_text": text,
        }
    if ext == ".docx":
        text = extract_docx(path)
        return {
            "pages": [{"page": 1, "text": text, "vision": False}],
            "method": "docx",
            "slides": [],
            "full_text": text,
        }
    if ext != ".pdf":
        return {"pages": [], "method": "", "slides": [], "full_text": ""}

    pages = await asyncio.to_thread(pdftotext_pages, path)
    slides_dir = dest / "slides"
    pngs: List[Path] = []
    is_slides = kind == "slides"
    need_vision = is_slides or any(len((t or "").strip()) < 20 for _, t in pages)
    if need_vision:
        pngs = await asyncio.to_thread(pdftoppm_png, path, str(slides_dir))

    max_vision = int(getattr(config, "COURSE_VISION_MAX_PAGES", 80) or 80)
    vision_used = 0
    out_pages = []
    method = "pdftotext"
    for i, (page_no, text) in enumerate(pages or [(n + 1, "") for n in range(len(pngs))]):
        t = (text or "").strip()
        used_vision = False
        if len(t) < 20 and vision_used < max_vision:
            png = None
            if pngs:
                if i < len(pngs):
                    png = pngs[i]
            if png and png.is_file() and openai_client:
                t = await vision_page(png, openai_client, user_id)
                used_vision = True
                vision_used += 1
                method = "vision" if method == "pdftotext" and not any(
                    p.get("text") for p in out_pages
                ) else "pdftotext+vision"
        out_pages.append({"page": page_no, "text": t, "vision": used_vision})

    full = "\n\n".join(
        f"[стр. {p['page']}]\n{p['text']}" for p in out_pages if p.get("text")
    )
    return {
        "pages": out_pages,
        "method": method,
        "slides": [str(p) for p in pngs],
        "full_text": full,
        "vision_pages": vision_used,
        "vision_limit_hit": vision_used >= max_vision and any(
            len((p.get("text") or "").strip()) < 20 and not p.get("vision")
            for p in out_pages
        ),
    }
